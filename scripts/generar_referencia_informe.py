"""Genera scripts/informe_referencia.docx: el documento de referencia de
pandoc con el formato del instructivo "El informe escrito" (v2.1,
Escuela de Ingeniería Industrial, FCEIA-UNR).

- A4, márgenes de 2,5 cm en los cuatro lados.
- Arial 10, interlineado simple, texto justificado y sin guiones.
- Encabezado con el título del trabajo; pie con los autores y
  "Página X de N" a la derecha.
- Cada parte de primer nivel (capítulos, prólogo, etc.) empieza en una
  página nueva.
- Estilos propios: "Carátula" (centrado) y "Figura" (imagen centrada).

Uso (python-docx no es dependencia del proyecto):
    uv run --with python-docx python scripts/generar_referencia_informe.py
"""

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TITULO = "Diseño de un Sistema de Información para la Asignación de Aulas"
AUTORES = "Luciano Antonelli – Pablo Galliano"
FUENTE = "Arial"
DESTINO = Path(__file__).with_name("informe_referencia.docx")


def _fuente(estilo, tam=None, negrita=None, cursiva=None):
    estilo.font.name = FUENTE
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FUENTE)
    if rfonts.getparent() is None:
        rpr.append(rfonts)
    if tam:
        estilo.font.size = Pt(tam)
    if negrita is not None:
        estilo.font.bold = negrita
    if cursiva is not None:
        estilo.font.italic = cursiva
    estilo.font.color.rgb = None


def _borde(parrafo, lado: str):
    """Línea fina arriba o abajo del párrafo (separador de cabecera/pie)."""
    ppr = parrafo._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{lado}")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "4"), ("w:color", "8FAADC")):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def _campo(parrafo, instruccion: str):
    """Inserta un campo de Word (PAGE, NUMPAGES) en el párrafo."""
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, instruccion), ("end", None)):
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {texto} "
            run._r.append(it)


def main():
    tmp = DESTINO.with_suffix(".base.docx")
    subprocess.run(["pandoc", "-o", str(tmp), "--print-default-data-file", "reference.docx"],
                   check=True)
    doc = Document(str(tmp))
    tmp.unlink()

    # Página.
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        for lado in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(s, lado, Cm(2.5))
        s.header_distance = s.footer_distance = Cm(1.25)

        # Cabecera: el título del trabajo, entre comillas y en gris, con
        # una línea debajo (ejemplo de la Figura 5 del instructivo).
        enc = s.header.paragraphs[0]
        enc.text = f"\u201c{TITULO}\u201d"
        enc.alignment = WD_ALIGN_PARAGRAPH.LEFT
        enc.runs[0].font.size = Pt(8)
        enc.runs[0].font.name = FUENTE
        enc.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        _borde(enc, "bottom")

        pie = s.footer.paragraphs[0]
        pie.text = ""
        pie.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
        # Pie: autores a la izquierda y "Página X de N" a la derecha, con
        # una línea arriba (ejemplo de la Figura 4 del instructivo).
        _borde(pie, "top")
        pie.add_run(f"Autores: {AUTORES}\tPágina ")
        _campo(pie, "PAGE")
        pie.add_run(" de ")
        _campo(pie, "NUMPAGES")
        for run in pie.runs:
            run.font.size = Pt(8)
            run.font.name = FUENTE
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    estilos = doc.styles
    por_nombre = {s.name: s for s in estilos}
    est = lambda n: por_nombre[n]
    # Cuerpo: Arial 10, simple, justificado.
    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if nombre in por_nombre:
            st = est(nombre)
            _fuente(st, 10)
            pf = st.paragraph_format
            pf.line_spacing = 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if nombre in ("Body Text", "First Paragraph"):
                pf.space_before, pf.space_after = Pt(0), Pt(6)
    # Títulos.
    for nivel, tam in ((1, 16), (2, 13), (3, 11), (4, 10)):
        st = est(f"Heading {nivel}")
        _fuente(st, tam, negrita=True, cursiva=False)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.space_before = Pt(18 if nivel == 1 else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        if nivel == 1:
            st.paragraph_format.page_break_before = True
    for nombre in ("Title", "Subtitle"):
        _fuente(est(nombre), 18 if nombre == "Title" else 13)
    # Estilos propios.
    car = estilos.add_style("Carátula", WD_STYLE_TYPE.PARAGRAPH)
    car.base_style = est("Normal")
    _fuente(car, 14)
    car.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    car.paragraph_format.space_after = Pt(18)
    fig = estilos.add_style("Figura", WD_STYLE_TYPE.PARAGRAPH)
    fig.base_style = est("Normal")
    fig.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.keep_with_next = True
    ley = estilos.add_style("Leyenda de figura", WD_STYLE_TYPE.PARAGRAPH)
    ley.base_style = est("Normal")
    _fuente(ley, 10)
    ley.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ley.paragraph_format.keep_with_next = True
    ley.paragraph_format.space_after = Pt(0)
    nota = estilos.add_style("Nota de figura", WD_STYLE_TYPE.PARAGRAPH)
    nota.base_style = est("Normal")
    _fuente(nota, 9)
    nota.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    nota.paragraph_format.space_after = Pt(10)

    # Tablas al estilo APA: línea arriba, debajo del encabezado y al pie,
    # sin líneas verticales; encabezado centrado.
    tabla = next(s for s in doc.styles if s.name == "Table")
    tpr = tabla.element.find(qn("w:tblPr"))
    if tpr is None:
        tpr = OxmlElement("w:tblPr"); tabla.element.append(tpr)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "bottom"):
        b = OxmlElement(f"w:{lado}")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "0"), ("w:color", "000000")):
            b.set(qn(k), v)
        bordes.append(b)
    tpr.append(bordes)
    for vieja in tabla.element.findall(qn("w:tblStylePr")):
        if vieja.get(qn("w:type")) == "firstRow":
            tabla.element.remove(vieja)
    fila = OxmlElement("w:tblStylePr"); fila.set(qn("w:type"), "firstRow")
    ppr = OxmlElement("w:pPr"); jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc)
    fila.append(ppr)
    tcpr = OxmlElement("w:tcPr"); tcb = OxmlElement("w:tcBorders")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "000000")):
        b.set(qn(k), v)
    tcb.append(b); tcpr.append(tcb); fila.append(tcpr)
    tabla.element.append(fila)

    # Sin guiones separadores de sílabas.
    settings = doc.settings.element
    for el in settings.findall(qn("w:autoHyphenation")):
        settings.remove(el)
    ah = OxmlElement("w:autoHyphenation")
    ah.set(qn("w:val"), "false")
    settings.append(ah)

    doc.save(str(DESTINO))
    print(f"✓ {DESTINO}")


if __name__ == "__main__":
    main()

"""Tests para template_export_service.

Fase C1 del rediseño 2026-09-15: verifica la generación de la
plantilla Excel de cronograma. Foco en la estructura correcta del
archivo generado (hojas presentes, DataValidation aplicado a las
columnas correctas, listas de códigos válidos) y en los guards
razonables (ciclo inexistente, ciclo sin dictados).
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine

from src.database.models import (
    CarreraDB,
    CicloDB,
    CicloPlanVersionDB,
    MateriaDB,
    PlanCarreraVersionDB,
    PlanEstudioDB,
)
from src.services.dictado_service import create_dictados_for_ciclo
from src.services.template_export_service import (
    generar_plantilla_cronograma_excel,
    obtener_referencia_materias_del_ciclo,
)


# =============================================================================
# Fixtures
# =============================================================================

@pytest.fixture(name="engine")
def engine_fixture():
    eng = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(eng)
    return eng


@pytest.fixture(name="session")
def session_fixture(engine):
    with Session(engine) as session:
        yield session


@pytest.fixture
def ciclo_con_2_materias(session):
    """Ciclo 2025-1C con MAT101 y FIS101 activos como dictados."""
    session.add(CarreraDB(codigo="ING", nombre="Ingeniería"))
    session.flush()

    pv = PlanCarreraVersionDB(
        id=str(uuid.uuid4()), carrera_codigo="ING",
        nombre="Plan A", fecha_creacion=date(2025, 1, 1),
    )
    session.add(pv)
    session.flush()

    session.add(MateriaDB(
        codigo="MAT101", nombre="Análisis I",
        periodo="cuatrimestral", active=True, horas_semanales=6,
    ))
    session.add(MateriaDB(
        codigo="FIS101", nombre="Física I",
        periodo="cuatrimestral", active=True, horas_semanales=6,
    ))
    session.flush()

    for mc in ("MAT101", "FIS101"):
        session.add(PlanEstudioDB(
            plan_version_id=pv.id, materia_codigo=mc,
            carrera_codigo="ING", anio_plan=1, cuatrimestre_plan="1C",
        ))

    ciclo = CicloDB(
        id="2025-1C", anio=2025, numero=1,
        fecha_inicio=date(2025, 3, 10), fecha_fin=date(2025, 7, 5),
    )
    session.add(ciclo)
    session.flush()
    session.add(CicloPlanVersionDB(ciclo_id=ciclo.id, plan_version_id=pv.id))
    session.commit()
    create_dictados_for_ciclo(session, ciclo.id)

    return {"ciclo": ciclo, "pv": pv}


# =============================================================================
# Tests
# =============================================================================


class TestGuards:
    def test_ciclo_inexistente_falla(self, session):
        with pytest.raises(ValueError, match="no existe"):
            generar_plantilla_cronograma_excel(session, "ciclo-fantasma")

    def test_ciclo_sin_dictados_falla(self, session, ciclo_con_2_materias):
        """Sin dictados creados, la plantilla no tiene lista de códigos
        posibles → mejor abortar con mensaje claro que generar un
        dropdown vacío.
        """
        # Crear un ciclo separado sin dictados
        c2 = CicloDB(
            id="2025-2C", anio=2025, numero=2,
            fecha_inicio=date(2025, 8, 1), fecha_fin=date(2025, 11, 30),
        )
        session.add(c2)
        session.commit()

        with pytest.raises(ValueError, match="dictados"):
            generar_plantilla_cronograma_excel(session, c2.id)


class TestEstructuraDelArchivo:
    def test_plantilla_tiene_hojas_esperadas(self, session, ciclo_con_2_materias):
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        assert "Instrucciones" in wb.sheetnames
        assert "Horarios" in wb.sheetnames
        # Hoja VISIBLE de referencia (2026-09-23): reemplaza a la
        # oculta `_materias` y expone codigo+nombre.
        assert "Materias" in wb.sheetnames
        assert "_materias" not in wb.sheetnames
        # Hojas ocultas de listas
        assert "_dias" in wb.sheetnames
        assert "_tipos" in wb.sheetnames
        assert "_virtual" in wb.sheetnames

    def test_hojas_de_listas_estan_ocultas(self, session, ciclo_con_2_materias):
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        for nombre in ("_dias", "_tipos", "_virtual"):
            assert wb[nombre].sheet_state == "hidden", (
                f"Hoja {nombre} debería estar oculta"
            )
        # La referencia de materias es VISIBLE (2026-09-23).
        assert wb["Materias"].sheet_state == "visible"

    def test_instrucciones_es_la_primera_hoja(self, session, ciclo_con_2_materias):
        """`Instrucciones` es la pestaña activa al abrir el archivo."""
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        assert wb.sheetnames[0] == "Instrucciones"
        assert wb.active.title == "Instrucciones"

    def test_headers_horarios_correctos(self, session, ciclo_con_2_materias):
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Horarios"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, 10)]
        # Esquema 2026-09-23: comisión con código numérico obligatorio
        # + nombre opcional, y materia elegible por nombre.
        assert headers == [
            "codigo_materia", "nombre_materia",
            "codigo_comision", "nombre_comision",
            "dia", "hora_inicio", "hora_fin",
            "tipo_clase", "virtual",
        ]

    def test_fila_2_esta_vacia_para_evitar_import_del_ejemplo(
        self, session, ciclo_con_2_materias,
    ):
        """Bugfix task #341 (2026-09-22): la fila 2 no debe contener el
        ejemplo `MAT101 · Lunes · 08:00 · 11:00`. Antes se escribía en
        amarillo como guía visual pero el parser no distingue ejemplo
        de dato real: si el usuario subía la plantilla sin borrar la
        fila 2, el ejemplo se importaba como horario válido.
        """
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Horarios"]
        # La columna A lleva la fórmula de auto-población del código
        # (2026-09-23) — que evalúa a "" mientras no se elija nombre.
        _a2 = ws.cell(row=2, column=1).value
        assert _a2 is None or str(_a2).startswith("=IFERROR"), (
            f"Fila 2 col A debería ser fórmula o vacía: {_a2!r}"
        )
        # Las demás celdas de la fila 2 deben estar vacías (sin datos
        # de ejemplo).
        for col in range(2, 10):
            assert ws.cell(row=2, column=col).value in (None, ""), (
                f"Fila 2 columna {col} no está vacía: "
                f"{ws.cell(row=2, column=col).value!r}"
            )

    def test_roundtrip_template_no_importa_filas_fantasma(
        self, session, ciclo_con_2_materias,
    ):
        """Regresión task #341: descargar la plantilla y pasarla al
        parser no debe devolver ningún ``HorarioInput`` — el usuario
        aún no cargó nada.
        """
        from src.services.horario_file_parser import parse_horarios_file

        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        upload = io.BytesIO(contenido)
        upload.name = "plantilla.xlsx"  # type: ignore[attr-defined]
        inputs, errors = parse_horarios_file(upload)
        assert inputs == [], (
            f"El parser debería devolver una lista vacía, obtuvo: {inputs}"
        )
        assert errors == [], f"No debería haber errores, obtuvo: {errors}"


class TestListasDeValores:
    def test_lista_de_materias_incluye_solo_dictados_activos(
        self, session, ciclo_con_2_materias,
    ):
        """La hoja `Materias` tiene exactamente los códigos con
        dictado activo en el ciclo (con su nombre al lado). Ni más
        ni menos.
        """
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Materias"]
        # Fila 1 = header ("codigo" / "nombre"); datos desde fila 2.
        codigos = [
            ws.cell(row=r, column=1).value
            for r in range(2, ws.max_row + 1)
        ]
        nombres = [
            ws.cell(row=r, column=2).value
            for r in range(2, ws.max_row + 1)
        ]
        assert sorted(codigos) == ["FIS101", "MAT101"]
        assert sorted(nombres) == ["Análisis I", "Física I"]

    def test_lista_dias_estandar(self, session, ciclo_con_2_materias):
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["_dias"]
        dias = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert dias == [
            "Lunes", "Martes", "Miércoles", "Jueves", "Viernes",
            "Sábado", "Domingo",
        ]

    def test_lista_tipos_clase(self, session, ciclo_con_2_materias):
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["_tipos"]
        tipos = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert tipos == ["teorica", "laboratorio"]

    def test_lista_virtual_si_no(self, session, ciclo_con_2_materias):
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["_virtual"]
        opts = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert opts == ["SI", "NO"]


class TestDataValidations:
    def test_hay_data_validation_en_columnas_esperadas(
        self, session, ciclo_con_2_materias,
    ):
        """La hoja Horarios tiene DataValidation configurados para
        codigo_materia (A), nombre_materia (B), codigo_comision (C),
        dia (E), hora_inicio (F), hora_fin (G), tipo_clase (H) y
        virtual (I). La columna D (nombre_comision) es texto libre y
        no lleva validación.
        """
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Horarios"]
        # openpyxl expone las DataValidation en ws.data_validations
        ranges_cubiertos: list[str] = []
        for dv in ws.data_validations.dataValidation:
            for r in dv.sqref.ranges:
                ranges_cubiertos.append(str(r))

        # Cada columna crítica debe tener al menos un rango.
        cols_letra = ["A", "B", "C", "E", "F", "G", "H", "I"]
        for letra in cols_letra:
            assert any(
                r.startswith(f"{letra}2") for r in ranges_cubiertos
            ), (
                f"Falta DataValidation para columna {letra}. "
                f"Rangos vistos: {ranges_cubiertos}"
            )

    def test_lista_materias_referencia_hoja_materias(
        self, session, ciclo_con_2_materias,
    ):
        """La validación de la columna A (codigo_materia) referencia la
        hoja visible `Materias` para armar la lista desplegable.
        """
        ciclo = ciclo_con_2_materias["ciclo"]
        contenido = generar_plantilla_cronograma_excel(session, ciclo.id)

        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Horarios"]

        dv_materias = None
        for dv in ws.data_validations.dataValidation:
            if any(str(r).startswith("A2") for r in dv.sqref.ranges):
                dv_materias = dv
                break

        assert dv_materias is not None
        assert dv_materias.type == "list"
        assert "Materias!" in (dv_materias.formula1 or "")


class TestReferenciaMaterias:
    def test_devuelve_pares_codigo_nombre(self, session, ciclo_con_2_materias):
        ciclo = ciclo_con_2_materias["ciclo"]
        refs = obtener_referencia_materias_del_ciclo(session, ciclo.id)
        assert refs == [
            ("FIS101", "Física I"),
            ("MAT101", "Análisis I"),
        ]

    def test_devuelve_vacio_si_no_hay_dictados(self, session):
        c = CicloDB(
            id="2025-2C", anio=2025, numero=2,
            fecha_inicio=date(2025, 8, 1), fecha_fin=date(2025, 11, 30),
        )
        session.add(c)
        session.commit()
        assert obtener_referencia_materias_del_ciclo(session, c.id) == []

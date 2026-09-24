"""Tests para inscripcion_import_service y plantilla de inscriptos.

Fase E1 del rediseño 2026-09-15. Cubre el preview + commit del
importer de inscriptos y la generación de la plantilla Excel
correspondiente.
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from src.database.models import (
    InscripcionHistoricaDB,
    MateriaDB,
)
from src.services.inscripcion_import_service import (
    InscripcionImportPreview,
    commit_import,
    preview_import,
    registrar_alias,
)
from src.services.template_export_service import (
    generar_plantilla_inscriptos_excel,
    obtener_referencia_materias_activas,
)


# =============================================================================
# Fixtures / helpers
# =============================================================================


@pytest.fixture(name="engine")
def engine_fixture():
    eng = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(eng)
    return eng


@pytest.fixture(name="session")
def session_fixture(engine):
    with Session(engine) as session:
        yield session


@pytest.fixture
def catalogo_basico(session):
    """Catálogo con 3 materias activas + 1 inactiva."""
    session.add(MateriaDB(
        codigo="MAT101", nombre="Análisis I",
        periodo="cuatrimestral", active=True, horas_semanales=6,
    ))
    session.add(MateriaDB(
        codigo="FIS101", nombre="Física I",
        periodo="cuatrimestral", active=True, horas_semanales=6,
    ))
    session.add(MateriaDB(
        codigo="QUI101", nombre="Química I",
        periodo="anual", active=True, horas_semanales=4,
    ))
    session.add(MateriaDB(
        codigo="LEGACY", nombre="Materia archivada",
        periodo="cuatrimestral", active=False, horas_semanales=3,
    ))
    session.commit()


def _fake_excel(df: pd.DataFrame, filename: str = "insc.xlsx") -> io.BytesIO:
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)
    buf.name = filename  # type: ignore[attr-defined]
    return buf


# =============================================================================
# Preview: parsing y validaciones
# =============================================================================


class TestPreviewParsing:
    def test_archivo_valido_arma_preview(self, session, catalogo_basico):
        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 120},
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "2C", "inscriptos": 90},
            {"codigo_materia": "QUI101", "anio": 2024,
             "cuatrimestre": "Anual", "inscriptos": 60},
        ])
        pv = preview_import(session, _fake_excel(df))

        assert pv.parse_errors == []
        assert pv.filas_error == []
        assert len(pv.filas_ok) == 3
        assert pv.n_nuevos == 3
        assert pv.n_pisan == 0

    def test_columnas_faltantes_bloquea(self, session, catalogo_basico):
        df = pd.DataFrame([{"codigo_materia": "MAT101", "anio": 2024}])
        pv = preview_import(session, _fake_excel(df))
        assert pv.tiene_errores_bloqueantes
        assert any("Columnas faltantes" in e for e in pv.parse_errors)

    def test_alias_de_columnas(self, session, catalogo_basico):
        """Los alias del loader legacy ('codigo', 'año', 'cuatri',
        'cant_inscriptos') se normalizan a los canónicos."""
        df = pd.DataFrame([
            {"codigo": "MAT101", "año": 2024,
             "cuatri": "1C", "cant_inscriptos": 100},
        ])
        pv = preview_import(session, _fake_excel(df))
        assert pv.parse_errors == []
        assert len(pv.filas_ok) == 1


class TestPreviewValidacionPorFila:
    def test_codigo_inexistente_va_a_error(self, session, catalogo_basico):
        df = pd.DataFrame([
            {"codigo_materia": "FANTASMA", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 50},
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 120},
        ])
        pv = preview_import(session, _fake_excel(df))
        assert len(pv.filas_ok) == 1
        assert len(pv.filas_error) == 1
        assert "FANTASMA" in pv.filas_error[0][1]

    def test_cuatrimestre_invalido_va_a_error(self, session, catalogo_basico):
        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "3C", "inscriptos": 50},
        ])
        pv = preview_import(session, _fake_excel(df))
        assert pv.filas_ok == []
        assert len(pv.filas_error) == 1
        assert "3C" in pv.filas_error[0][1] or "cuatrimestre" in pv.filas_error[0][1]

    def test_cuatri_case_insensitive(self, session, catalogo_basico):
        """'anual' se normaliza a 'Anual', '1c' a '1C'."""
        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "anual", "inscriptos": 50},
            {"codigo_materia": "FIS101", "anio": 2024,
             "cuatrimestre": "1c", "inscriptos": 30},
        ])
        pv = preview_import(session, _fake_excel(df))
        assert pv.filas_error == []
        assert {f.cuatrimestre for f in pv.filas_ok} == {"Anual", "1C"}

    def test_inscriptos_negativo_va_a_error(self, session, catalogo_basico):
        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": -5},
        ])
        pv = preview_import(session, _fake_excel(df))
        assert pv.filas_ok == []
        assert len(pv.filas_error) == 1
        assert "negativo" in pv.filas_error[0][1]

    def test_anio_fuera_de_rango(self, session, catalogo_basico):
        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 1985,
             "cuatrimestre": "1C", "inscriptos": 50},
        ])
        pv = preview_import(session, _fake_excel(df))
        assert pv.filas_ok == []
        assert len(pv.filas_error) == 1


class TestPreviewConValoresPrevios:
    def test_marca_filas_que_pisan(self, session, catalogo_basico):
        """Una fila cuyo (materia, año, cuatri) ya existe se marca con
        valor_previo poblado y cambia_valor=True si el valor difiere.
        """
        session.add(InscripcionHistoricaDB(
            materia_codigo="MAT101", anio=2023,
            cuatrimestre="1C", inscriptos=100,
        ))
        session.commit()

        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2023,
             "cuatrimestre": "1C", "inscriptos": 150},
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 200},
        ])
        pv = preview_import(session, _fake_excel(df))

        fila_pisa = next(f for f in pv.filas_ok if f.anio == 2023)
        fila_nueva = next(f for f in pv.filas_ok if f.anio == 2024)

        assert fila_pisa.valor_previo == 100
        assert fila_pisa.cambia_valor is True
        assert fila_nueva.valor_previo is None
        assert fila_nueva.es_nuevo is True

        assert pv.n_nuevos == 1
        assert pv.n_pisan == 1

    def test_marca_filas_sin_cambio(self, session, catalogo_basico):
        """Fila con valor idéntico al previo → es_nuevo=False,
        cambia_valor=False, cuenta en n_iguales."""
        session.add(InscripcionHistoricaDB(
            materia_codigo="MAT101", anio=2023,
            cuatrimestre="1C", inscriptos=100,
        ))
        session.commit()

        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2023,
             "cuatrimestre": "1C", "inscriptos": 100},
        ])
        pv = preview_import(session, _fake_excel(df))
        assert pv.n_iguales == 1
        assert pv.n_pisan == 0
        assert pv.n_nuevos == 0


class TestPreviewDuplicadosEnArchivo:
    def test_duplicado_ultima_fila_gana(self, session, catalogo_basico):
        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 100},
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 150},
        ])
        pv = preview_import(session, _fake_excel(df))

        assert len(pv.filas_ok) == 1
        assert pv.filas_ok[0].inscriptos == 150
        assert any("duplica" in w for w in pv.warnings)


# =============================================================================
# Commit
# =============================================================================


class TestCommit:
    def test_commit_crea_filas_nuevas(self, session, catalogo_basico):
        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 120},
            {"codigo_materia": "FIS101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 80},
        ])
        pv = preview_import(session, _fake_excel(df))
        result = commit_import(session, pv)

        assert result.filas_creadas == 2
        assert result.filas_actualizadas == 0
        assert result.filas_sin_cambio == 0

        rows = session.exec(select(InscripcionHistoricaDB)).all()
        assert len(rows) == 2

    def test_commit_actualiza_valores_existentes(
        self, session, catalogo_basico,
    ):
        session.add(InscripcionHistoricaDB(
            materia_codigo="MAT101", anio=2023,
            cuatrimestre="1C", inscriptos=100,
        ))
        session.commit()

        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2023,
             "cuatrimestre": "1C", "inscriptos": 150},
        ])
        pv = preview_import(session, _fake_excel(df))
        result = commit_import(session, pv)

        assert result.filas_actualizadas == 1
        assert result.filas_creadas == 0

        db = session.exec(select(InscripcionHistoricaDB)).one()
        assert db.inscriptos == 150

    def test_commit_sin_cambio_no_incrementa_actualizadas(
        self, session, catalogo_basico,
    ):
        session.add(InscripcionHistoricaDB(
            materia_codigo="MAT101", anio=2023,
            cuatrimestre="1C", inscriptos=100,
        ))
        session.commit()

        df = pd.DataFrame([
            {"codigo_materia": "MAT101", "anio": 2023,
             "cuatrimestre": "1C", "inscriptos": 100},
        ])
        pv = preview_import(session, _fake_excel(df))
        result = commit_import(session, pv)

        assert result.filas_actualizadas == 0
        assert result.filas_sin_cambio == 1

    def test_commit_ignora_filas_error(self, session, catalogo_basico):
        """Filas con error no bloquean el commit del resto."""
        df = pd.DataFrame([
            {"codigo_materia": "FANTASMA", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 50},
            {"codigo_materia": "MAT101", "anio": 2024,
             "cuatrimestre": "1C", "inscriptos": 120},
        ])
        pv = preview_import(session, _fake_excel(df))
        result = commit_import(session, pv)

        assert result.filas_creadas == 1

    def test_commit_bloqueado_por_errores_estructurales(self, session):
        pv = InscripcionImportPreview()
        pv.parse_errors.append("Columnas faltantes")
        with pytest.raises(ValueError, match="bloqueantes"):
            commit_import(session, pv)


# =============================================================================
# Plantilla Excel
# =============================================================================


class TestPlantillaInscriptos:
    def test_genera_archivo_con_hojas_esperadas(self, session, catalogo_basico):
        contenido = generar_plantilla_inscriptos_excel(session)

        wb = load_workbook(io.BytesIO(contenido))
        assert "Instrucciones" in wb.sheetnames
        assert "Inscriptos" in wb.sheetnames
        # 2026-09-24: hoja VISIBLE `Materias` con contexto (reemplaza
        # a la oculta `_materias`), como en la plantilla de cronograma.
        assert "Materias" in wb.sheetnames
        assert "_materias" not in wb.sheetnames
        assert "_cuatris" in wb.sheetnames

    def test_hoja_materias_con_contexto_y_protegida(
        self, session, catalogo_basico,
    ):
        """La hoja `Materias` es visible, protegida (sólo consulta) y
        trae el contexto del catálogo — incluido el código Guaraní,
        que las cátedras suelen necesitar para cruzar sus planillas.
        Sólo materias activas (LEGACY no aparece).
        """
        contenido = generar_plantilla_inscriptos_excel(session)
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Materias"]
        assert ws.protection.sheet is True
        headers = [
            ws.cell(row=1, column=c).value
            for c in range(1, ws.max_column + 1)
        ]
        for esperado in (
            "Nombre", "Código", "Código Guaraní", "Período",
            "Optativa", "Virtual (catálogo)", "Planes de carrera",
        ):
            assert any(
                str(h).startswith(esperado) for h in headers if h
            ), f"Falta header {esperado!r}: {headers}"
        codigos = [
            ws.cell(row=r, column=2).value
            for r in range(2, ws.max_row + 1)
        ]
        assert sorted(codigos) == ["FIS101", "MAT101", "QUI101"]

    def test_lista_cuatris(self, session, catalogo_basico):
        contenido = generar_plantilla_inscriptos_excel(session)
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["_cuatris"]
        vals = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
        assert vals == ["1C", "2C", "Anual"]

    def test_headers_correctos(self, session, catalogo_basico):
        contenido = generar_plantilla_inscriptos_excel(session)
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Inscriptos"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, 6)]
        # 2026-09-24: espejo de la plantilla de cronograma — nombre
        # primero (fórmula de sólo lectura) y código como única
        # entrada de materia.
        assert headers == [
            "nombre_materia", "codigo_materia", "anio",
            "cuatrimestre", "inscriptos",
        ]

    def test_nombre_formula_y_hoja_protegida(self, session, catalogo_basico):
        contenido = generar_plantilla_inscriptos_excel(session)
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Inscriptos"]
        _a2 = str(ws.cell(row=2, column=1).value or "")
        assert _a2.startswith("=IFERROR")
        assert "Materias!$A$" in _a2 and "$B2" in _a2
        assert ws.protection.sheet is True
        assert ws.cell(row=2, column=1).protection.locked is True
        for col in range(2, 6):
            assert ws.cell(row=2, column=col).protection.locked is False

    def test_hoja_inscriptos_es_tabla(self, session, catalogo_basico):
        contenido = generar_plantilla_inscriptos_excel(session)
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Inscriptos"]
        assert "TablaInscriptos" in ws.tables

    def test_roundtrip_plantilla_sin_filas_fantasma(
        self, session, catalogo_basico,
    ):
        """Subir la plantilla recién generada no debe producir filas
        ni errores: las filas con sólo la fórmula del nombre se
        saltean en silencio (mismo comportamiento que la plantilla
        de cronograma).
        """
        contenido = generar_plantilla_inscriptos_excel(session)
        upload = io.BytesIO(contenido)
        upload.name = "plantilla_inscriptos.xlsx"
        pv = preview_import(session, upload)
        assert pv.parse_errors == []
        assert pv.filas_error == []
        assert pv.filas_ok == []

    def test_datavalidation_configurada(self, session, catalogo_basico):
        """Columnas de carga con DataValidation: código (B), año (C),
        cuatrimestre (D), inscriptos (E). El nombre (A) es fórmula y
        no lleva validación.
        """
        contenido = generar_plantilla_inscriptos_excel(session)
        wb = load_workbook(io.BytesIO(contenido))
        ws = wb["Inscriptos"]
        cols_cubiertas: set[str] = set()
        for dv in ws.data_validations.dataValidation:
            for r in dv.sqref.ranges:
                col_letra = str(r).split(":")[0][0]
                cols_cubiertas.add(col_letra)
        assert {"B", "C", "D", "E"}.issubset(cols_cubiertas)
        assert "A" not in cols_cubiertas

    def test_sin_materias_activas_falla(self, session):
        with pytest.raises(ValueError, match="No hay materias"):
            generar_plantilla_inscriptos_excel(session)

    def test_referencia_activas_helper(self, session, catalogo_basico):
        refs = obtener_referencia_materias_activas(session)
        codigos = [c for c, _ in refs]
        assert "LEGACY" not in codigos
        assert "MAT101" in codigos


class TestCodigosNoResueltos:
    """2026-09-24: la vista previa expone los códigos sin match para
    que la UI ofrezca la asociación (alias) sin salir del flujo — el
    importador fuerza que todo código matchee o se revise ahí mismo.
    """

    def test_preview_lista_codigos_sin_match(self, session, catalogo_basico):
        archivo = io.BytesIO(
            b"codigo_materia,anio,cuatrimestre,inscriptos\n"
            b"NOEXISTE,2024,1C,50\n"
            b"NOEXISTE,2023,1C,40\n"
            b"OTRORARO,2024,2C,10\n"
            b"MAT101,2024,1C,80\n"
        )
        archivo.name = "insc.csv"
        pv = preview_import(session, archivo)
        assert pv.codigos_no_resueltos == ["NOEXISTE", "OTRORARO"]
        assert len(pv.filas_error) == 3
        assert len(pv.filas_ok) == 1

    def test_alias_registrado_resuelve_en_siguiente_preview(
        self, session, catalogo_basico,
    ):
        registrar_alias(session, "NOEXISTE", "MAT101", origen="manual")
        archivo = io.BytesIO(
            b"codigo_materia,anio,cuatrimestre,inscriptos\n"
            b"NOEXISTE,2024,1C,50\n"
        )
        archivo.name = "insc.csv"
        pv = preview_import(session, archivo)
        assert pv.codigos_no_resueltos == []
        assert pv.filas_error == []
        assert len(pv.filas_ok) == 1
        assert pv.filas_ok[0].materia_codigo == "MAT101"

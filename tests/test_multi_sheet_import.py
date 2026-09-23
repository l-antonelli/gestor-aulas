"""Tests para el selector de hoja del importer (Fase I.4, 2026-09-23).

Cover:
- ``list_horarios_sheets`` y ``list_inscriptos_sheets`` reconocen las
  hojas visibles y esconden las de sistema (``Instrucciones``, ``_*``).
- ``parse_horarios_file(sheet_name=X)`` respeta la hoja pedida y
  cae al fallback tradicional cuando ``sheet_name`` es ``None``.
- ``preview_import`` (cronograma e inscriptos) propaga
  ``sheet_name`` al parser.
"""

from __future__ import annotations

import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine

from src.database.models import (
    CarreraDB,
    CicloDB,
    CicloPlanVersionDB,
    MateriaDB,
    PlanCarreraVersionDB,
    PlanEstudioDB,
    ScheduleDB,
)


@pytest.fixture(name="engine")
def engine_fixture():
    eng = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(eng)
    return eng


@pytest.fixture(name="session")
def session_fixture(engine):
    with Session(engine) as s:
        yield s


def _excel_con_hojas(sheets: dict[str, pd.DataFrame], filename: str) -> io.BytesIO:
    """Crea un Excel en memoria con múltiples hojas."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf) as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
    buf.seek(0)
    buf.name = filename  # type: ignore[attr-defined]
    return buf


def _excel_con_hoja_oculta(
    sheets_visible: dict[str, pd.DataFrame],
    sheets_hidden: dict[str, list[str]],
    filename: str,
) -> io.BytesIO:
    """Crea un Excel donde algunas hojas quedan `sheet_state='hidden'`."""
    wb = Workbook()
    # Eliminar la hoja default.
    default = wb.active
    if default is not None:
        wb.remove(default)
    for name, df in sheets_visible.items():
        ws = wb.create_sheet(title=name)
        ws.append(list(df.columns))
        for _, row in df.iterrows():
            ws.append([row[c] for c in df.columns])
    for name, rows in sheets_hidden.items():
        ws = wb.create_sheet(title=name)
        ws.sheet_state = "hidden"
        for r in rows:
            ws.append([r])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = filename  # type: ignore[attr-defined]
    return buf


class TestListHorariosSheets:
    def test_devuelve_hojas_visibles(self):
        from src.services.horario_file_parser import list_horarios_sheets

        excel = _excel_con_hojas(
            {
                "1C": pd.DataFrame([{"codigo_materia": "M", "dia": "Lunes",
                                     "hora_inicio": "8:00", "hora_fin": "10:00"}]),
                "2C": pd.DataFrame([{"codigo_materia": "N", "dia": "Martes",
                                     "hora_inicio": "9:00", "hora_fin": "11:00"}]),
            },
            "cronogramas.xlsx",
        )
        assert list_horarios_sheets(excel) == ["1C", "2C"]

    def test_excluye_instrucciones_y_hojas_de_sistema(self):
        from src.services.horario_file_parser import list_horarios_sheets

        excel = _excel_con_hoja_oculta(
            sheets_visible={
                "Instrucciones": pd.DataFrame([{"guía": "leer esto"}]),
                "Horarios": pd.DataFrame([{"codigo_materia": "M", "dia": "Lunes",
                                           "hora_inicio": "8:00", "hora_fin": "10:00"}]),
            },
            sheets_hidden={
                "_materias": ["M"],
                "_dias": ["Lunes"],
            },
            filename="template.xlsx",
        )
        assert list_horarios_sheets(excel) == ["Horarios"]

    def test_csv_devuelve_vacio(self):
        from src.services.horario_file_parser import list_horarios_sheets

        buf = io.BytesIO(b"codigo_materia,dia,hora_inicio,hora_fin\n"
                         b"M,Lunes,08:00,10:00\n")
        buf.name = "cronograma.csv"  # type: ignore[attr-defined]
        assert list_horarios_sheets(buf) == []


class TestParseHorariosFileConSheetName:
    def test_sheet_name_explicito_gana(self):
        """Si el usuario elige `2C`, el parser lee esa hoja aunque
        también exista `Horarios` con contenido distinto.
        """
        from src.services.horario_file_parser import parse_horarios_file

        excel = _excel_con_hojas(
            {
                "Horarios": pd.DataFrame([
                    {"codigo_materia": "MAT101", "dia": "Lunes",
                     "hora_inicio": "8:00", "hora_fin": "10:00"},
                ]),
                "2C": pd.DataFrame([
                    {"codigo_materia": "FIS201", "dia": "Miércoles",
                     "hora_inicio": "14:00", "hora_fin": "16:00"},
                ]),
            },
            "cats.xlsx",
        )
        entries, errs = parse_horarios_file(excel, sheet_name="2C")
        assert errs == []
        assert len(entries) == 1
        assert entries[0].codigo_materia == "FIS201"

    def test_sheet_name_none_usa_fallback(self):
        """Sin `sheet_name`, el parser prefiere `Horarios`."""
        from src.services.horario_file_parser import parse_horarios_file

        excel = _excel_con_hojas(
            {
                "Horarios": pd.DataFrame([
                    {"codigo_materia": "MAT101", "dia": "Lunes",
                     "hora_inicio": "8:00", "hora_fin": "10:00"},
                ]),
                "2C": pd.DataFrame([
                    {"codigo_materia": "FIS201", "dia": "Miércoles",
                     "hora_inicio": "14:00", "hora_fin": "16:00"},
                ]),
            },
            "cats.xlsx",
        )
        entries, errs = parse_horarios_file(excel)
        assert errs == []
        assert len(entries) == 1
        assert entries[0].codigo_materia == "MAT101"

    def test_sheet_name_inexistente_cae_al_fallback(self):
        """Si el sheet_name pedido no existe, se cae al fallback (no
        se levanta excepción — el usuario ya no podrá elegir esa hoja
        vía UI, pero por robustez del service).
        """
        from src.services.horario_file_parser import parse_horarios_file

        excel = _excel_con_hojas(
            {
                "Horarios": pd.DataFrame([
                    {"codigo_materia": "MAT101", "dia": "Lunes",
                     "hora_inicio": "8:00", "hora_fin": "10:00"},
                ]),
            },
            "cats.xlsx",
        )
        entries, errs = parse_horarios_file(excel, sheet_name="NO_EXISTE")
        assert errs == []
        assert len(entries) == 1
        assert entries[0].codigo_materia == "MAT101"


class TestPreviewImportCronogramaConSheetName:
    def _catalog(self, session):
        session.add(CarreraDB(codigo="I", nombre="Informática"))
        session.flush()
        pv = PlanCarreraVersionDB(
            id="pv1", carrera_codigo="I", nombre="P",
            fecha_creacion=date(2025, 1, 1),
        )
        session.add(pv)
        session.add(MateriaDB(
            codigo="MAT101", nombre="An", periodo="cuatrimestral",
            active=True, horas_semanales=4,
        ))
        session.add(MateriaDB(
            codigo="FIS201", nombre="Fis", periodo="cuatrimestral",
            active=True, horas_semanales=4,
        ))
        session.flush()
        for mc in ("MAT101", "FIS201"):
            session.add(PlanEstudioDB(
                plan_version_id="pv1", materia_codigo=mc,
                carrera_codigo="I", anio_plan=1, cuatrimestre_plan="1C",
            ))
        ciclo = CicloDB(
            id="2025-1C", anio=2025, numero=1,
            fecha_inicio=date(2025, 3, 1), fecha_fin=date(2025, 7, 1),
        )
        session.add(ciclo)
        session.flush()
        session.add(CicloPlanVersionDB(ciclo_id="2025-1C", plan_version_id="pv1"))
        sched = ScheduleDB(
            id="sched1", ciclo_id="2025-1C", nombre="test",
            fecha_upload=date(2025, 3, 1),
        )
        session.add(sched)
        session.commit()
        return sched

    def test_preview_lee_la_hoja_elegida(self, session):
        """El caller elige `2C` y `preview_import` procesa esa hoja."""
        from src.services.cronograma_import_service import preview_import

        sched = self._catalog(session)
        excel = _excel_con_hojas(
            {
                "1C": pd.DataFrame([
                    {"codigo_materia": "MAT101", "dia": "Lunes",
                     "hora_inicio": "8:00", "hora_fin": "10:00"},
                ]),
                "2C": pd.DataFrame([
                    {"codigo_materia": "FIS201", "dia": "Martes",
                     "hora_inicio": "14:00", "hora_fin": "16:00"},
                ]),
            },
            "cats.xlsx",
        )

        pv = preview_import(session, sched.id, excel, sheet_name="2C")
        assert pv.parse_errors == []
        cods = {m.materia_codigo for m in pv.materias}
        assert cods == {"FIS201"}


class TestListInscriptosSheets:
    def test_devuelve_hojas_visibles(self):
        from src.services.inscripcion_import_service import list_inscriptos_sheets

        excel = _excel_con_hojas(
            {
                "2024": pd.DataFrame([{"codigo_materia": "M", "anio": 2024,
                                       "cuatrimestre": "1C", "inscriptos": 100}]),
                "2025": pd.DataFrame([{"codigo_materia": "M", "anio": 2025,
                                       "cuatrimestre": "1C", "inscriptos": 120}]),
            },
            "inscriptos.xlsx",
        )
        assert list_inscriptos_sheets(excel) == ["2024", "2025"]

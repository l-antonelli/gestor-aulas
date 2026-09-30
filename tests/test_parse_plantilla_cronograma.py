"""Tests de `parse_plantilla_cronograma` (2026-09-30).

El importador del flujo "crear cronograma con archivo" deja de tomar
una sola hoja: lee la plantilla multihoja completa que genera la app
(una hoja por grupo de materias) y la acepta sólo si es exactamente
esa plantilla (hoja oculta `_meta` del mismo ciclo, mismas hojas, sin
materias repetidas entre hojas). Así se fuerza el uso de la plantilla
tal como sale de la aplicación.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from src.services.horario_file_parser import parse_plantilla_cronograma
from src.services.template_export_service import (
    HOJA_META,
    generar_plantilla_cronograma_excel,
)
from tests.test_template_export_service import (  # noqa: F401 (fixtures)
    ciclo_con_2_materias,
    ciclo_con_grupos,
    engine_fixture,
    session_fixture,
)

HOJA_MAT = "Básicas"                 # MAT101
HOJA_FIS = "FísicaEspeciales rara"   # FIS101


def _fila(ws, r, codigo, com, dia, hi, hf):
    for c, v in zip(range(2, 8), (codigo, com, f"C{com}", dia, hi, hf)):
        ws.cell(row=r, column=c, value=v)


def _archivo(wb) -> io.BytesIO:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "plantilla.xlsx"
    return buf


def _plantilla(session, ciclo_id="2025-1C"):
    return load_workbook(io.BytesIO(
        generar_plantilla_cronograma_excel(session, ciclo_id),
    ))


def test_plantilla_valida_lee_todas_las_hojas(session, ciclo_con_grupos):
    wb = _plantilla(session)
    _fila(wb[HOJA_MAT], 2, "MAT101", 1, "Lunes", "08:00", "11:00")
    _fila(wb[HOJA_MAT], 3, "MAT101", 1, "Miércoles", "08:00", "11:00")
    _fila(wb[HOJA_FIS], 2, "FIS101", 1, "Martes", "14:00", "17:00")

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2025-1C")

    assert res.errores == []
    assert len(res.entries) == 3
    assert res.materia_a_hoja == {"MAT101": HOJA_MAT, "FIS101": HOJA_FIS}
    por_hoja = {h.nombre: h for h in res.hojas}
    assert [h.nombre for h in res.hojas] == [HOJA_MAT, HOJA_FIS]
    assert por_hoja[HOJA_MAT].grupo == "Básicas"
    assert por_hoja[HOJA_MAT].codigos == ["MAT101"]
    assert por_hoja[HOJA_MAT].n_horarios == 2
    assert por_hoja[HOJA_FIS].n_horarios == 1


def test_hoja_vacia_no_es_error(session, ciclo_con_grupos):
    wb = _plantilla(session)
    _fila(wb[HOJA_MAT], 2, "MAT101", 1, "Lunes", "08:00", "11:00")

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2025-1C")

    assert res.errores == []
    assert {h.nombre: h.n_horarios for h in res.hojas} == {
        HOJA_MAT: 1, HOJA_FIS: 0,
    }


def test_archivo_sin_meta_se_rechaza():
    wb = Workbook()
    ws = wb.active
    ws.title = "Horarios"
    for c, h in enumerate(("codigo_materia", "dia", "hora_inicio", "hora_fin"), 1):
        ws.cell(row=1, column=c, value=h)
    ws.append(["MAT101", "Lunes", "08:00", "11:00"])

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2025-1C")

    assert res.entries == []
    assert len(res.errores) == 1
    assert "plantilla" in res.errores[0]
    assert "descarg" in res.errores[0].lower()


def test_plantilla_de_otro_ciclo_se_rechaza(session, ciclo_con_grupos):
    wb = _plantilla(session)
    _fila(wb[HOJA_MAT], 2, "MAT101", 1, "Lunes", "08:00", "11:00")

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2026-2C")

    assert res.entries == []
    assert "2025-1C" in res.errores[0] and "2026-2C" in res.errores[0]


def test_version_vieja_se_rechaza(session, ciclo_con_grupos):
    wb = _plantilla(session)
    ws = wb[HOJA_META]
    for row in ws.iter_rows():
        if row[0].value == "version":
            row[1].value = 1

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2025-1C")

    assert res.entries == []
    assert "versión" in res.errores[0]


def test_hoja_renombrada_se_rechaza(session, ciclo_con_grupos):
    wb = _plantilla(session)
    wb[HOJA_FIS].title = "Mi hoja"

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2025-1C")

    assert res.entries == []
    texto = " ".join(res.errores)
    assert f"'{HOJA_FIS}'" in texto   # falta la original
    assert "'Mi hoja'" in texto       # sobra la nueva


def test_materia_en_dos_hojas_es_bloqueante(session, ciclo_con_grupos):
    wb = _plantilla(session)
    _fila(wb[HOJA_MAT], 2, "MAT101", 1, "Lunes", "08:00", "11:00")
    _fila(wb[HOJA_FIS], 2, "MAT101", 2, "Martes", "08:00", "11:00")

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2025-1C")

    (err,) = [e for e in res.errores if "MAT101" in e]
    assert f"'{HOJA_MAT}'" in err and f"'{HOJA_FIS}'" in err


def test_error_de_fila_indica_la_hoja(session, ciclo_con_grupos):
    wb = _plantilla(session)
    _fila(wb[HOJA_FIS], 2, "FIS101", 1, "Martes", "17:00", "14:00")

    res = parse_plantilla_cronograma(_archivo(wb), ciclo_id_esperado="2025-1C")

    assert res.errores
    assert all(e.startswith(f"Hoja '{HOJA_FIS}'") for e in res.errores)
    (hoja_fis,) = [h for h in res.hojas if h.nombre == HOJA_FIS]
    assert hoja_fis.errores


def test_sin_ciclo_esperado_no_chequea_el_ciclo(session, ciclo_con_grupos):
    wb = _plantilla(session)
    res = parse_plantilla_cronograma(_archivo(wb))
    assert res.errores == []
    assert res.meta is not None and res.meta.ciclo_id == "2025-1C"
